#!/usr/bin/env python3
"""
Figure 2 worked examples as NATIVE PowerPoint shapes.

Every node is a real rounded rectangle with editable text and every is_a edge is
a real connector with an arrowhead, so the whole figure can be edited directly in
PowerPoint (no flattened images). Styling matches Figure 1 of the ESID4HPO deck.

Slide size follows the deck used as TEMPLATE (10 x 7.5 in).

Structures were extracted from the real releases (v2024-08-13 vs v2026-06-23);
re-verify before submission, e.g.
    grep "HP_5210283" hp-edit.owl | grep SubClassOf

Deps:  pip install python-pptx lxml
Run:   python fig2_build_native_pptx.py
"""

import pathlib

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = pathlib.Path(__file__).resolve().parent
TEMPLATE = HERE / "esid4hpo_figures.pptx"          # deck whose theme/size to reuse
OUT = HERE / "_work" / "figures" / "ESID4HPO_Figure2_editable.pptx"

NAVY = RGBColor(0x00, 0x37, 0x54)
BLUE = RGBColor(0x4E, 0x7E, 0x96)
GREY = RGBColor(0xE7, 0xEA, 0xEC)
AMBER = RGBColor(0xC8, 0xA8, 0x70)
TEAL = RGBColor(0x6C, 0xBC, 0xC5)
INK = RGBColor(0x2B, 0x2F, 0x33)
MUT = RGBColor(0x5F, 0x70, 0x78)
BORD = RGBColor(0xC4, 0xC9, 0xCE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ARROW = RGBColor(0x9A, 0xA3, 0xAD)

# kind -> (fill, line, line_pt, text_colour, bold, dashed)
STYLE = {
    "root":  (NAVY, NAVY, 1.0, WHITE, True, False),
    "new":   (BLUE, BLUE, 1.0, WHITE, True, False),
    "rep":   (GREY, AMBER, 2.25, INK, False, False),
    "relab": (GREY, TEAL, 2.25, INK, False, False),
    "exist": (GREY, BORD, 1.0, INK, False, False),
    "ctx":   (WHITE, BORD, 1.0, MUT, False, True),
}

BOX_H, VGAP, HGAP = 0.42, 0.30, 0.09


def drop_all_slides(prs):
    lst = prs.slides._sldIdLst
    for s in list(lst):
        rId = s.get("{http://schemas.openxmlformats.org/officeDocument/2006/"
                    "relationships}id")
        prs.part.drop_rel(rId)
        lst.remove(s)


def white_bg(slide):
    """BIH master maps bg1 -> dk1 (navy); figure panels need white."""
    cSld = slide._element.find(qn("p:cSld"))
    for old in cSld.findall(qn("p:bg")):
        cSld.remove(old)
    bg = etree.SubElement(cSld, qn("p:bg"))
    bgPr = etree.SubElement(bg, qn("p:bgPr"))
    sf = etree.SubElement(bgPr, qn("a:solidFill"))
    c = etree.SubElement(sf, qn("a:srgbClr"))
    c.set("val", "FFFFFF")
    etree.SubElement(bgPr, qn("a:effectLst"))
    cSld.insert(0, bg)


def textbox(s, x, y, w, h, text, size=9, bold=False, color=INK,
            align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return tb


def box(s, x, y, w, h, label, kind, size):
    fill, line, lw, tc, bold, dashed = STYLE[kind]
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                            Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    if dashed:
        ln = sh.line._get_or_add_ln()
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "dash")
    try:
        sh.adjustments[0] = 0.13          # softer corner radius
    except Exception:
        pass
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = label
    p.alignment = PP_ALIGN.CENTER
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = tc
    return sh


def arrow(s, x1, y1, x2, y2, color=ARROW, width=1.0):
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                Inches(x2), Inches(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(width)
    ln = cn.line._get_or_add_ln()
    te = etree.SubElement(ln, qn("a:tailEnd"))
    te.set("type", "triangle")
    te.set("w", "med")
    te.set("len", "med")
    return cn


def draw_panel(s, panel, x0, w_panel, y0):
    """panel = {'rows': [[(id,label,kind), ...], ...],
                'edges': [(child, parent[, 'amber'|'teal']), ...]}"""
    pos = {}
    for i, row in enumerate(panel["rows"]):
        n = len(row)
        bw = (w_panel - HGAP * (n - 1)) / n
        size = 8.0 if bw > 1.5 else (7.0 if bw > 1.15 else 6.0)
        y = y0 + i * (BOX_H + VGAP)
        for j, (nid, label, kind) in enumerate(row):
            x = x0 + j * (bw + HGAP)
            box(s, x, y, bw, BOX_H, label, kind, size)
            pos[nid] = (x + bw / 2, y, y + BOX_H)
    for e in panel["edges"]:
        child, parent = e[0], e[1]
        st = e[2] if len(e) > 2 else None
        if child not in pos or parent not in pos:
            continue
        cx, ctop, _ = pos[child]
        px, _, pbot = pos[parent]
        col = AMBER if st == "amber" else (TEAL if st == "teal" else ARROW)
        wdt = 1.75 if st in ("amber", "teal") else 1.0
        arrow(s, cx, ctop, px, pbot, col, wdt)
    return pos


def add_example(prs, layout, letter, title, before, after, caption):
    s = prs.slides.add_slide(layout)
    white_bg(s)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)

    frame = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.18), Inches(0.08),
                               Inches(9.64), Inches(7.10))
    frame.fill.background()
    frame.line.color.rgb = BORD
    frame.line.width = Pt(0.75)
    frame.shadow.inherit = False

    textbox(s, 0.24, 0.12, 0.5, 0.30, letter, size=16, bold=True, color=MUT)
    textbox(s, 0.72, 0.14, 8.9, 0.34, title, size=13, bold=True, color=NAVY)
    textbox(s, 0.30, 0.56, 4.45, 0.26, "Before   (v2024-08-13)", size=10,
            bold=True, color=MUT, align=PP_ALIGN.CENTER)
    textbox(s, 5.25, 0.56, 4.45, 0.26, "After   (v2026-06-23)", size=10,
            bold=True, color=MUT, align=PP_ALIGN.CENTER)

    draw_panel(s, before, 0.30, 4.45, 0.92)
    draw_panel(s, after, 5.25, 4.45, 0.92)

    lx, ly = 0.34, 5.62
    for fill, line, txt in [(BLUE, BLUE, "new term"), (GREY, AMBER, "re-parented"),
                            (GREY, TEAL, "relabelled"), (GREY, BORD, "unchanged"),
                            (WHITE, BORD, "context")]:
        sq = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(lx), Inches(ly),
                                Inches(0.19), Inches(0.19))
        sq.fill.solid()
        sq.fill.fore_color.rgb = fill
        sq.line.color.rgb = line
        sq.line.width = Pt(1.5)
        sq.shadow.inherit = False
        textbox(s, lx + 0.25, ly + 0.005, 1.5, 0.2, txt, size=9, color=INK)
        lx += 0.25 + 0.075 * len(txt) + 0.22
    textbox(s, 0.34, 5.98, 9.25, 1.05, caption, size=9, color=INK)
    return s


# ============================== EXAMPLES ==================================
FLOW_B = {"rows": [
    [("TC", "Abnormal T cell count", "ctx")],
    [("TS", "Abnormal T cell subset distribution", "exist")],
    [("P4", "Abnormal proportion of CD4-positive T cells", "exist")],
    [("D4", "Decreased proportion of CD4-positive T cells", "exist"),
     ("I4", "Increased proportion of CD4-positive T cells", "exist")]],
    "edges": [("TS", "TC"), ("P4", "TS"), ("D4", "P4"), ("I4", "P4")]}

FLOW_A = {"rows": [
    [("TM", "Abnormal T cell morphology", "ctx")],
    [("TS", "Abnormal T cell subset number", "relab")],
    [("N4", "Abnormal total CD4+ T cell number", "new")],
    [("DN", "Decreased total CD4+ T cell number", "new"),
     ("IN", "Increased total CD4+ T cell number", "new")],
    [("DC", "Decreased total CD4+ T cell count", "new"),
     ("DP", "Decreased total CD4+ T cell proportion", "relab"),
     ("IC", "Increased total CD4+ T cell count", "new"),
     ("IP", "Increased total CD4+ T cell proportion", "relab")]],
    "edges": [("TS", "TM", "amber"), ("N4", "TS"), ("DN", "N4"), ("IN", "N4"),
              ("DC", "DN"), ("DP", "DN", "amber"), ("IC", "IN"),
              ("IP", "IN", "amber")]}

AB_B = {"rows": [
    [("CONC", "Decreased circulating antibody concentration", "exist"),
     ("AGR", "Impaired antigen-specific response", "ctx")],
    [("SP", "Decreased circulating level of specific antibody", "exist")],
    [("POLY", "Decreased specific anti-polysaccharide antibody level", "exist"),
     ("PNEU", "Decreased specific pneumococcal antibody level", "exist")]],
    "edges": [("SP", "CONC"), ("SP", "AGR"), ("POLY", "SP"), ("PNEU", "SP")]}

AB_A = {"rows": [
    [("HUM", "Abnormality of humoral immunity", "ctx"),
     ("AGR", "Impaired antigen-specific response", "ctx")],
    [("SP", "Impaired specific antibody response", "relab"),
     ("VAC", "Decreased response to unconjugated polysaccharide vaccine", "exist")],
    [("POLY", "Decreased specific anti-polysaccharide antibody concentration", "relab"),
     ("TY1", "Complete absence of response to S. typhi vaccine", "new"),
     ("TY2", "Partial absence of response to S. typhi vaccine", "new")],
    [("PNEU", "Decreased specific pneumococcal antibody concentration", "rep"),
     ("HIB", "Decreased specific Haemophilus influenzae b antibody concentration", "new")]],
    "edges": [("SP", "HUM", "amber"), ("SP", "AGR"), ("POLY", "SP"),
              ("TY1", "VAC"), ("TY2", "VAC"), ("PNEU", "POLY", "amber"),
              ("HIB", "POLY")]}

INF_B = {"rows": [
    [("UI", "Unusual infection", "exist"),
     ("RTI", "Respiratory tract infection", "ctx")],
    [("SITE", "Unusual infection by anatomical site", "exist"),
     ("VIR", "Unusual viral infection", "exist"),
     ("PNE", "Pneumonia", "exist")],
    [("CNS", "Unusual CNS infection", "exist"),
     ("GI", "Unusual gastrointestinal infection", "exist"),
     ("SKIN", "Unusual skin infection", "exist")]],
    "edges": [("SITE", "UI"), ("VIR", "UI"), ("PNE", "RTI"),
              ("CNS", "SITE"), ("GI", "SITE"), ("SKIN", "SITE")]}

INF_A = {"rows": [
    [("UI", "Unusual infection", "exist")],
    [("SITE", "Unusual infection by anatomical site", "exist"),
     ("VIR", "Unusual viral infection", "exist"),
     ("OPP", "Opportunistic viral infection", "exist")],
    [("LRTI", "Unusual lower respiratory tract infection", "new"),
     ("CMV", "Unusual cytomegalovirus infection", "new")],
    [("PNE", "Pneumonia", "rep"),
     ("CMVP", "Cytomegalovirus pneumonitis", "new")]],
    "edges": [("SITE", "UI"), ("VIR", "UI"), ("LRTI", "SITE"), ("CMV", "VIR"),
              ("PNE", "LRTI", "amber"), ("CMVP", "LRTI"), ("CMVP", "CMV"),
              ("CMVP", "OPP")]}

GRAN_B = {"rows": [
    [("MAC", "Abnormal macrophage morphology", "ctx"),
     ("PI", "Abnormal pulmonary interstitial morphology", "ctx"),
     ("GRAN", "Granuloma", "ctx")],
    [("G", "Granulomatosis", "exist"),
     ("PG", "Pulmonary granulomatosis", "exist"),
     ("EG", "Eosinophilic granuloma", "exist")],
    [("NPG", "Necrotizing pulmonary granulomatosis", "exist"),
     ("NNPG", "Non-necrotizing pulmonary granulomatosis", "exist")]],
    "edges": [("G", "MAC"), ("PG", "PI"), ("EG", "GRAN"),
              ("NPG", "PG"), ("NNPG", "PG")]}

GRAN_A = {"rows": [
    [("G", "Granulomatosis", "root")],
    [("INF", "Infectious granulomatosis", "new"),
     ("NINF", "Non-infectious granulomatosis", "new"),
     ("GLN", "Granulomatous lymphadenopathy", "new")],
    [("CAS", "Caseating granulomatosis", "new"),
     ("NCAS", "Non-caseating granulomatosis", "new"),
     ("GA", "Granuloma annulare", "new"),
     ("EG", "Eosinophilic granuloma", "rep")],
    [("NPG", "Necrotizing pulmonary granulomatosis", "rep"),
     ("NNPG", "Non-necrotizing pulmonary granulomatosis", "rep")],
    [("GLILD", "Granulomatous lymphocytic interstitial lung disease (GLILD)", "new")]],
    "edges": [("INF", "G"), ("NINF", "G"), ("GLN", "G"), ("CAS", "INF"),
              ("NCAS", "NINF"), ("GA", "NINF"), ("EG", "NINF", "amber"),
              ("NPG", "CAS", "amber"), ("NNPG", "NCAS", "amber"),
              ("GLILD", "NNPG")]}

SLIDES = [
    ("a)", "Flow cytometry: separating absolute counts from proportions",
     FLOW_B, FLOW_A,
     "Before curation, CD4+ T cell subsets could only be recorded as proportions and no term existed for an "
     "absolute count. After curation an etiology-free \u201cnumber\u201d parent was introduced for each direction of "
     "change, with absolute count and proportion as distinct children (HP:5210418 vs HP:0032218). Existing "
     "proportion terms were relabelled from laboratory nomenclature to clinical wording and re-parented accordingly."),
    ("b)", "Antibodies: response separated from immunoglobulin concentration",
     AB_B, AB_A,
     "\u201cDecreased circulating level of specific antibody\u201d was classified as a subtype of decreased antibody "
     "concentration. It was relabelled \u201cImpaired specific antibody response\u201d (HP:0012475) and re-parented under "
     "humoral immunity, since an impaired response is not a subtype of a reduced concentration. New terms cover "
     "Haemophilus influenzae type b and graded responses to unconjugated Salmonella typhi vaccine."),
    ("c)", "Infections: cross-classification by organism, site and opportunism",
     INF_B, INF_A,
     "The anatomical-site axis previously held only three classes and \u201cPneumonia\u201d was not reachable from the "
     "unusual-infection branch. Eight further site classes were added and organ-specific infections are now "
     "cross-classified on three axes: HP:5210283 Cytomegalovirus pneumonitis is a child of Opportunistic viral "
     "infection, Unusual cytomegalovirus infection and Unusual lower respiratory tract infection, so the case is "
     "retrieved by a query on any axis."),
    ("d)", "Granulomatosis: an etiological axis for granulomatous disease",
     GRAN_B, GRAN_A,
     "Granulomatosis terms previously hung off unrelated morphological parents with no etiological organisation. "
     "An infectious / non-infectious \u00d7 caseating / non-caseating scaffold was created under Granulomatosis, "
     "existing terms were re-parented into it, and clinically important entities were added, including "
     "granulomatous lymphocytic interstitial lung disease (GLILD)."),
]


def main():
    if not TEMPLATE.exists():
        raise SystemExit(f"{TEMPLATE} not found - put the deck next to this script.")
    prs = Presentation(str(TEMPLATE))
    drop_all_slides(prs)
    layout = next((l for l in prs.slide_layouts if l.name == "Bilder zweispaltig"),
                  prs.slide_layouts[0])
    for letter, title, b, a, cap in SLIDES:
        add_example(prs, layout, letter, title, b, a, cap)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"saved {OUT}  ({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
